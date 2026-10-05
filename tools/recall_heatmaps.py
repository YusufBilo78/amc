"""Recall per class at every SNR, in the layout of the comparison table Yusuf
was shown (2016 vs 2018 at 18 / 0 / -6 dB, then one class x SNR grid per
dataset), built from our own runs.

    python tools/recall_heatmaps.py       -> recall_tables.xlsx, figures/36_*.png, figures/37_*.png

Sources: the 11-class RML2016.10a run (train_backbone_rml2016_f1000_colab.npz,
3 seeds x 150 test frames = 450 decisions per cell) and the 24-class RadioML
2018 run (train_backbone_rml2018_f512_colab.npz, 3 seeds x 76-77 test frames
per cell). Recall = diagonal / row, seeds pooled, in percent. A dataset whose
file has no `confusions_by_snr` yet is reported as pending, not guessed.

The 2018 run used 512 of the 4,096 frames per cell, so a 2018 cell rests on
about 230 decisions and a 2016 cell on 450; one decision is 0.4 and 0.2
points respectively.
"""
import pathlib

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
RUNS = {"2016": "train_backbone_rml2016_f1000_colab.npz",
        "2018": "train_backbone_rml2018_f512_colab.npz"}
# 2016 class -> its counterpart in 2018 (the comparison table's rows)
PAIRS = [("BPSK", "BPSK"), ("QPSK", "QPSK"), ("8PSK", "8PSK"), ("QAM16", "16QAM"),
         ("QAM64", "64QAM"), ("PAM4", "4ASK"), ("WBFM", "FM"), ("AM-DSB", "AM-DSB-SC")]
LEVELS = [18, 0, -6]
CMAP = LinearSegmentedColormap.from_list("rg", ["#E8705A", "#F5D76E", "#5BC85B"])
HEAD = PatternFill("solid", fgColor="2F4A6D")


def load(key):
    z = np.load(ROOT / RUNS[key], allow_pickle=True)
    if "confusions_by_snr" not in z.files:
        return None
    by = z["confusions_by_snr"].sum(0)                     # snr, n, n
    rec = 100 * np.diagonal(by, axis1=1, axis2=2) / by.sum(2)
    return {"classes": [str(c) for c in z["class_names"]], "snrs": z["snrs"].astype(int).tolist(),
            "recall": rec.T, "per_cell": int(by[0].sum(1).min()), "seeds": len(z["overall"])}


def heatmap(d, title, out):
    cls, snrs, r = d["classes"], d["snrs"], d["recall"]
    fig, ax = plt.subplots(figsize=(0.55 * len(snrs) + 2.2, 0.36 * len(cls) + 1.4))
    ax.imshow(r, cmap=CMAP, vmin=0, vmax=100, aspect="auto")
    for i in range(len(cls)):
        for j in range(len(snrs)):
            ax.text(j, i, f"{r[i, j]:.0f}", ha="center", va="center", fontsize=8, color="#1E293B")
    ax.set_xticks(range(len(snrs)), snrs, fontsize=8)
    ax.set_yticks(range(len(cls)), cls, fontsize=9)
    ax.xaxis.tick_top()
    ax.set_title(title, loc="left", fontsize=11, pad=22)
    for s in ax.spines.values():
        s.set_visible(False)
    fig.tight_layout()
    fig.savefig(out, dpi=160)
    plt.close(fig)


def write_grid(ws, d, title):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"recall %, {d['seeds']} seeds pooled, at least {d['per_cell']} test decisions per cell"])
    ws.append(["class \\ SNR (dB)"] + d["snrs"])
    for c in ws[3]:
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, Alignment(horizontal="center")
    for name, row in zip(d["classes"], d["recall"]):
        ws.append([name] + [round(float(v)) for v in row])
    last = get_column_letter(len(d["snrs"]) + 1)
    rng = f"B4:{last}{3 + len(d['classes'])}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="E8705A",
                                                      mid_type="num", mid_value=50, mid_color="F5D76E",
                                                      end_type="num", end_value=100, end_color="5BC85B"))
    ws.column_dimensions["A"].width = 16
    for k in range(2, len(d["snrs"]) + 2):
        ws.column_dimensions[get_column_letter(k)].width = 5.5


def main():
    data = {k: load(k) for k in RUNS}
    wb = Workbook()
    ws = wb.active
    ws.title = "2016 vs 2018"
    ws.append(["Class (2016 → 2018)"] + [f"2016 · {l} dB" for l in LEVELS] + [f"2018 · {l} dB" for l in LEVELS])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for a, b in PAIRS:
        row = [f"{a} → {b}"]
        for key, name in (("2016", a), ("2018", b)):
            d = data[key]
            for l in LEVELS:
                if d is None:
                    row.append("pending")
                else:
                    row.append(round(float(d["recall"][d["classes"].index(name), d["snrs"].index(l)])))
        ws.append(row)
    ws.conditional_formatting.add(f"B2:G{1 + len(PAIRS)}", ColorScaleRule(
        start_type="num", start_value=0, start_color="E8705A", mid_type="num", mid_value=50,
        mid_color="F5D76E", end_type="num", end_value=100, end_color="5BC85B"))
    ws.column_dimensions["A"].width = 24
    note = ["", "Recall % at that SNR label, seeds pooled. 2016: 11-class ICRNNA run, 450 decisions per cell.",
            "2018: 24-class ICRNNA run, about 230 decisions per cell." if data["2018"] else
            "2018: pending — per-SNR matrices are being rebuilt from the checkpoints.",
            "The SNR labels are not comparable between the datasets, nor between classes within 2016 (NOISE_2016.md)."]
    for t in note:
        ws.append([t])
    for key, title in (("2016", "RML2016.10a — 11 classes × 20 SNR"), ("2018", "RadioML 2018.01A — 24 classes × 26 SNR")):
        d = data[key]
        if d is None:
            wb.create_sheet(key).append([f"{title}: pending"])
            continue
        write_grid(wb.create_sheet(key), d, title)
        heatmap(d, f"{title} — recall %, {d['seeds']} seeds, ≥{d['per_cell']} decisions per cell",
                ROOT / "figures" / f"{36 if key == '2016' else 37}_recall_grid_{key}.png")
    wb.save(ROOT / "recall_tables.xlsx")
    for key, d in data.items():
        print(key, "pending" if d is None else f"{len(d['classes'])} x {len(d['snrs'])}, >= {d['per_cell']} per cell")


if __name__ == "__main__":
    main()
