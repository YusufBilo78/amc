"""Every modulation at every SNR, as one workbook.

Reads result files written by train_backbone.py that carry
`confusions_by_snr` and writes an .xlsx with, per run:

    <run> recall      rows = transmitted class, columns = SNR, cells = recall
                      in percent, computed by formula from the matrices sheet;
                      last row = accuracy over all classes at that SNR
    <run> matrices    the full confusion matrix at every SNR, counts, seeds
                      pooled, one block per level

and a Markdown file with the recall tables, for reading on a phone.

    python tools/snr_tables.py train_backbone_rml2016_f1000_c4_colab_s14.npz \
        train_backbone_rml2018_f2048_c4_colab_e150.npz -o snr_tables.xlsx

Counts are data; every percentage is a formula, so the workbook recalculates
if a count is edited.
"""
import argparse
import pathlib

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="DCE6F0")
DIAG = PatternFill("solid", fgColor="E6F4F1")


def col(c):
    return get_column_letter(c)


def sheet_name(stem, suffix):
    base = stem.replace("train_backbone_", "").replace("_colab", "")
    return (base + " " + suffix)[:31]


def write_run(wb, path, md):
    z = np.load(path, allow_pickle=True)
    names = [str(c) for c in z["class_names"]]
    snrs = z["snrs"].astype(int).tolist()
    by = z["confusions_by_snr"].sum(0)          # (n_snr, n, n), seeds pooled
    n, n_seeds = len(names), len(z["overall"])
    per_cell = int(by[0].sum(1)[0])
    stem = path.stem

    # ---------------------------------------------------------- matrices
    ms = wb.create_sheet(sheet_name(stem, "matrices"))
    ms["A1"] = f"{stem}: confusion matrix at every SNR, counts, {n_seeds} seeds pooled"
    ms["A1"].font = Font(name=FONT, bold=True, size=12)
    ms["A2"] = (f"rows = transmitted, columns = decided; every row sums to "
                f"{per_cell} decisions ({per_cell // n_seeds} test frames per (class, SNR) cell x {n_seeds} seeds)")
    ms["A2"].font = Font(name=FONT, italic=True, color="5C7A99")
    anchors = {}                                  # snr -> first data row
    r = 4
    for k, s in enumerate(snrs):
        ms.cell(r, 1, f"{s:+d} dB").font = Font(name=FONT, bold=True)
        ms.cell(r, 2, "transmitted \\ decided").font = Font(name=FONT, bold=True)
        for j, c in enumerate(names):
            cell = ms.cell(r, 3 + j, c); cell.font = Font(name=FONT, bold=True); cell.fill = HEAD
            cell.alignment = Alignment(horizontal="center")
        ms.cell(r, 3 + n, "row total").font = Font(name=FONT, bold=True)
        anchors[s] = r + 1
        for i, c in enumerate(names):
            ms.cell(r + 1 + i, 2, c).font = Font(name=FONT, bold=True)
            for j in range(n):
                cell = ms.cell(r + 1 + i, 3 + j, int(by[k, i, j]))
                cell.font = Font(name=FONT); cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="center")
                if i == j:
                    cell.fill = DIAG
            tot = ms.cell(r + 1 + i, 3 + n, f"=SUM({col(3)}{r + 1 + i}:{col(2 + n)}{r + 1 + i})")
            tot.font = Font(name=FONT, color="5C7A99"); tot.number_format = "#,##0"
        r += n + 3
    ms.column_dimensions["A"].width = 9
    ms.column_dimensions["B"].width = 22
    for j in range(n + 1):
        ms.column_dimensions[col(3 + j)].width = 11
    ms.freeze_panes = "C4"

    # ------------------------------------------------------------ recall
    rs = wb.create_sheet(sheet_name(stem, "recall"))
    rs["A1"] = f"{stem}: recall of every class at every SNR, percent"
    rs["A1"].font = Font(name=FONT, bold=True, size=12)
    rs["A2"] = (f"{per_cell} decisions behind every cell; the last row pools the {n} classes, "
                f"{per_cell * n} decisions per level. Chance is {100 / n:.1f}. "
                f"Formulas read the '{ms.title}' sheet.")
    rs["A2"].font = Font(name=FONT, italic=True, color="5C7A99")
    rs.cell(4, 1, "class \\ SNR (dB)").font = Font(name=FONT, bold=True)
    for k, s in enumerate(snrs):
        cell = rs.cell(4, 2 + k, s); cell.font = Font(name=FONT, bold=True); cell.fill = HEAD
        cell.alignment = Alignment(horizontal="center"); cell.number_format = "+0;-0;0"
    mt = f"'{ms.title}'"
    for i, c in enumerate(names):
        rs.cell(5 + i, 1, c).font = Font(name=FONT, bold=True)
        for k, s in enumerate(snrs):
            a = anchors[s]
            f = f"=100*{mt}!{col(3 + i)}{a + i}/{mt}!{col(3 + n)}{a + i}"
            cell = rs.cell(5 + i, 2 + k, f); cell.font = Font(name=FONT)
            cell.number_format = "0.0"; cell.alignment = Alignment(horizontal="center")
    row_all = 5 + n
    rs.cell(row_all, 1, f"all {n} classes").font = Font(name=FONT, bold=True)
    for k, s in enumerate(snrs):
        a = anchors[s]
        diag = "+".join(f"{mt}!{col(3 + i)}{a + i}" for i in range(n))
        f = f"=100*({diag})/SUM({mt}!{col(3 + n)}{a}:{col(3 + n)}{a + n - 1})"
        cell = rs.cell(row_all, 2 + k, f); cell.font = Font(name=FONT, bold=True)
        cell.number_format = "0.0"; cell.alignment = Alignment(horizontal="center")
    rs.column_dimensions["A"].width = 18
    for k in range(len(snrs)):
        rs.column_dimensions[col(2 + k)].width = 7
    rs.freeze_panes = "B5"

    # ---------------------------------------------------------- markdown
    md.append(f"## {stem}\n")
    md.append(f"Recall in percent; {per_cell} decisions per cell, {n_seeds} seeds pooled; "
              f"last row pools the {n} classes ({per_cell * n} per level); chance {100 / n:.1f}.\n")
    md.append("| class \\ SNR | " + " | ".join(f"{s:+d}" for s in snrs) + " |")
    md.append("|---|" + "---:|" * len(snrs))
    for i, c in enumerate(names):
        rec = [100 * by[k, i, i] / by[k, i].sum() for k in range(len(snrs))]
        md.append(f"| **{c}** | " + " | ".join(f"{v:.1f}" for v in rec) + " |")
    allacc = [100 * np.trace(by[k]) / by[k].sum() for k in range(len(snrs))]
    md.append(f"| *all {n}* | " + " | ".join(f"*{v:.1f}*" for v in allacc) + " |\n")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("npz", nargs="+")
    p.add_argument("-o", "--out", default="snr_tables.xlsx")
    args = p.parse_args()
    wb = Workbook(); wb.remove(wb.active)
    md = ["# Every class at every SNR\n"]
    for f in args.npz:
        write_run(wb, pathlib.Path(f), md)
    # recall sheets first, matrices after them, each group in run order
    wb._sheets = ([ws for ws in wb._sheets if ws.title.endswith(" recall")]
                  + [ws for ws in wb._sheets if ws.title.endswith(" matrices")])
    wb.active = 0
    out = pathlib.Path(args.out); wb.save(out)
    out.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")
    print("wrote", out, "and", out.with_suffix(".md"))


if __name__ == "__main__":
    main()
